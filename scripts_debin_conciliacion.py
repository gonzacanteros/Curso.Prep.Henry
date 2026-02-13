#!/usr/bin/env python3
"""Conciliación DEBIN entre archivos COELSA (.txt) y Core bancario (.csv).

Genera un Excel con resumen y detalle de la conciliación.
"""

from __future__ import annotations

import argparse
import csv
import ftplib
import re
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None

HEADER_CANDIDATES = {
    "debin_id": ["iddebin", "id_debin", "debinid", "debin", "id debin"],
    "concepto": ["concepto", "detalle", "descripcion", "descripción", "glosa"],
    "comnro": ["comnro", "comprobante", "nrocomprobante", "nro_comprobante", "numero comprobante"],
    "cuenta": ["cuenta", "cbu", "alias", "cta", "cuenta_destino", "cuenta_origen"],
    "importe": ["importe", "monto", "amount", "valor"],
}

DEBIN_REGEX = re.compile(r"\b(?:ID\s*DEBIN|DEBIN)\s*[:\-#]?\s*([A-Z0-9]{6,40})\b", re.IGNORECASE)


@dataclass
class Movement:
    source: str
    raw: Dict[str, str]
    debin_id: str
    comnro: str
    cuenta: str
    importe: Decimal

    @property
    def compound_key(self) -> Tuple[str, str, Decimal]:
        return (self.comnro, self.cuenta, self.importe)


@dataclass
class Totals:
    total_importe: Decimal
    creditos_importe: Decimal
    debitos_importe_abs: Decimal
    creditos_cantidad: int
    debitos_cantidad: int
    cantidad_total: int


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def normalize_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def parse_decimal(value: str) -> Decimal:
    txt = (value or "").strip().replace(" ", "")
    if not txt:
        return Decimal("0")

    if txt.count(",") > 0 and txt.count(".") > 0:
        txt = txt.replace(".", "").replace(",", ".")
    elif txt.count(",") > 0 and txt.count(".") == 0:
        txt = txt.replace(",", ".")

    try:
        return Decimal(txt).quantize(Decimal("0.01"))
    except InvalidOperation:
        return Decimal("0.00")


def resolve_column(headers: Sequence[str], candidates: Sequence[str]) -> Optional[str]:
    normalized_map = {normalize_key(h): h for h in headers}
    for candidate in candidates:
        key = normalize_key(candidate)
        if key in normalized_map:
            return normalized_map[key]
    return None


def detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,|\t").delimiter
    except csv.Error:
        for delim in [";", "|", "\t", ","]:
            if delim in sample:
                return delim
        return ";"


def read_tabular(path: Path) -> List[Dict[str, str]]:
    content = path.read_text(encoding="utf-8", errors="replace")
    sample = "\n".join(content.splitlines()[:20])
    delimiter = detect_delimiter(sample)

    rows: List[Dict[str, str]] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        if reader.fieldnames:
            for row in reader:
                rows.append({k or "": normalize_text(v or "") for k, v in row.items()})
        else:
            f.seek(0)
            lines = [normalize_text(x) for x in f.readlines() if normalize_text(x)]
            if not lines:
                return []
            headers = [f"col_{i+1}" for i in range(max(len(l.split(delimiter)) for l in lines))]
            for line in lines:
                parts = [normalize_text(x) for x in line.split(delimiter)]
                rows.append({headers[i]: parts[i] if i < len(parts) else "" for i in range(len(headers))})

    return rows


def get_value(row: Dict[str, str], key_name: str, fallback: str = "") -> str:
    headers = list(row.keys())
    col = resolve_column(headers, HEADER_CANDIDATES[key_name])
    if col:
        return normalize_text(row.get(col, ""))
    return fallback


def extract_debin_id(row: Dict[str, str]) -> str:
    direct = get_value(row, "debin_id")
    if direct:
        return direct.upper()

    concept = get_value(row, "concepto")
    match = DEBIN_REGEX.search(concept)
    if match:
        return match.group(1).upper()

    generic = re.search(r"\b([A-Z0-9]{10,40})\b", concept.upper())
    return generic.group(1) if generic else ""


def to_movement(row: Dict[str, str], source: str) -> Movement:
    concepto = get_value(row, "concepto")
    debin_id = extract_debin_id(row)
    comnro = re.sub(r"\s+", "", get_value(row, "comnro")).upper()
    cuenta = re.sub(r"\s+", "", get_value(row, "cuenta")).upper()
    importe = parse_decimal(get_value(row, "importe"))

    row_copy = dict(row)
    row_copy.setdefault("concepto_resuelto", concepto)

    return Movement(source=source, raw=row_copy, debin_id=debin_id, comnro=comnro, cuenta=cuenta, importe=importe)


def load_movements(paths: Iterable[Path], source: str) -> List[Movement]:
    out: List[Movement] = []
    for p in paths:
        for row in read_tabular(p):
            out.append(to_movement(row, source=source))
    return out


def reconcile(coelsa: List[Movement], core: List[Movement]) -> List[Dict[str, str]]:
    core_by_id: Dict[str, List[int]] = {}
    core_by_compound: Dict[Tuple[str, str, Decimal], List[int]] = {}
    core_used: set[int] = set()

    for idx, m in enumerate(core):
        if m.debin_id:
            core_by_id.setdefault(m.debin_id, []).append(idx)
        core_by_compound.setdefault(m.compound_key, []).append(idx)

    report: List[Dict[str, str]] = []

    for c in coelsa:
        row = {
            "status": "SOLO_COELSA",
            "coelsa_debin_id": c.debin_id,
            "core_debin_id": "",
            "coelsa_comnro": c.comnro,
            "core_comnro": "",
            "coelsa_cuenta": c.cuenta,
            "core_cuenta": "",
            "coelsa_importe": f"{c.importe:.2f}",
            "core_importe": "",
        }

        matched_idx: Optional[int] = None
        if c.debin_id:
            for idx in core_by_id.get(c.debin_id, []):
                if idx not in core_used:
                    matched_idx = idx
                    row["status"] = "MATCH_ID_DEBIN"
                    break

        if matched_idx is None:
            for idx in core_by_compound.get(c.compound_key, []):
                if idx not in core_used:
                    matched_idx = idx
                    row["status"] = "MATCH_COMNRO_CUENTA_IMPORTE"
                    break

        if matched_idx is not None:
            core_used.add(matched_idx)
            m = core[matched_idx]
            row.update(
                {
                    "core_debin_id": m.debin_id,
                    "core_comnro": m.comnro,
                    "core_cuenta": m.cuenta,
                    "core_importe": f"{m.importe:.2f}",
                }
            )

        report.append(row)

    for idx, m in enumerate(core):
        if idx in core_used:
            continue
        report.append(
            {
                "status": "SOLO_CORE",
                "coelsa_debin_id": "",
                "core_debin_id": m.debin_id,
                "coelsa_comnro": "",
                "core_comnro": m.comnro,
                "coelsa_cuenta": "",
                "core_cuenta": m.cuenta,
                "coelsa_importe": "",
                "core_importe": f"{m.importe:.2f}",
            }
        )

    return report


def calculate_totals(movs: List[Movement]) -> Totals:
    creditos = [m.importe for m in movs if m.importe > 0]
    debitos = [m.importe for m in movs if m.importe < 0]
    total = sum((m.importe for m in movs), Decimal("0.00"))
    return Totals(
        total_importe=total,
        creditos_importe=sum(creditos, Decimal("0.00")),
        debitos_importe_abs=abs(sum(debitos, Decimal("0.00"))),
        creditos_cantidad=len(creditos),
        debitos_cantidad=len(debitos),
        cantidad_total=len(movs),
    )


def build_summary_rows(coelsa: List[Movement], core: List[Movement], report: List[Dict[str, str]]) -> List[List[object]]:
    t_coelsa = calculate_totals(coelsa)
    t_core = calculate_totals(core)
    counts = Counter(r["status"] for r in report)

    def diff(a: Decimal, b: Decimal) -> Decimal:
        return (a - b).quantize(Decimal("0.01"))

    return [
        ["Métrica", "COELSA", "CORE", "Diferencia (COELSA - CORE)"],
        ["Importe total", float(t_coelsa.total_importe), float(t_core.total_importe), float(diff(t_coelsa.total_importe, t_core.total_importe))],
        ["Importe créditos", float(t_coelsa.creditos_importe), float(t_core.creditos_importe), float(diff(t_coelsa.creditos_importe, t_core.creditos_importe))],
        ["Importe débitos (abs)", float(t_coelsa.debitos_importe_abs), float(t_core.debitos_importe_abs), float(diff(t_coelsa.debitos_importe_abs, t_core.debitos_importe_abs))],
        ["Cantidad total", t_coelsa.cantidad_total, t_core.cantidad_total, t_coelsa.cantidad_total - t_core.cantidad_total],
        ["Cantidad créditos", t_coelsa.creditos_cantidad, t_core.creditos_cantidad, t_coelsa.creditos_cantidad - t_core.creditos_cantidad],
        ["Cantidad débitos", t_coelsa.debitos_cantidad, t_core.debitos_cantidad, t_coelsa.debitos_cantidad - t_core.debitos_cantidad],
        [],
        ["Estado conciliación", "Cantidad"],
        ["MATCH_ID_DEBIN", counts.get("MATCH_ID_DEBIN", 0)],
        ["MATCH_COMNRO_CUENTA_IMPORTE", counts.get("MATCH_COMNRO_CUENTA_IMPORTE", 0)],
        ["SOLO_COELSA (faltan en CORE)", counts.get("SOLO_COELSA", 0)],
        ["SOLO_CORE (sobran en CORE)", counts.get("SOLO_CORE", 0)],
    ]


def write_excel_report(path: Path, summary_rows: List[List[object]], report_rows: List[Dict[str, str]]) -> None:
    if Workbook is None:
        raise SystemExit("Falta dependencia openpyxl. Instalar con: pip install openpyxl")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Resumen"
    for row in summary_rows:
        ws_summary.append(row)

    ws_detail = wb.create_sheet("Detalle Conciliacion")
    headers = [
        "status",
        "coelsa_debin_id",
        "core_debin_id",
        "coelsa_comnro",
        "core_comnro",
        "coelsa_cuenta",
        "core_cuenta",
        "coelsa_importe",
        "core_importe",
    ]
    ws_detail.append(headers)
    for r in report_rows:
        ws_detail.append([r[h] for h in headers])

    ws_missing = wb.create_sheet("Faltan_en_CORE")
    ws_missing.append(headers)
    for r in report_rows:
        if r["status"] == "SOLO_COELSA":
            ws_missing.append([r[h] for h in headers])

    ws_surplus = wb.create_sheet("Sobran_en_CORE")
    ws_surplus.append(headers)
    for r in report_rows:
        if r["status"] == "SOLO_CORE":
            ws_surplus.append([r[h] for h in headers])

    wb.save(path)


def download_from_ftp(host: str, username: str, password: str, remote_dir: str, output_dir: Path, pattern: str) -> List[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    downloaded: List[Path] = []
    regex = re.compile("^" + pattern.replace(".", "\\.").replace("*", ".*") + "$")

    with ftplib.FTP(host) as ftp:
        ftp.login(username, password)
        ftp.cwd(remote_dir)
        for name in ftp.nlst():
            if not regex.match(name):
                continue
            local_path = output_dir / Path(name).name
            with local_path.open("wb") as f:
                ftp.retrbinary(f"RETR {name}", f.write)
            downloaded.append(local_path)

    return downloaded


def main() -> int:
    parser = argparse.ArgumentParser(description="Concilia movimientos DEBIN entre COELSA y Core")
    parser.add_argument("--core-csv", required=True, help="CSV de movimientos DEBIN del Core")
    parser.add_argument("--local-input-dir", default="./inputs", help="Directorio local con .txt de COELSA")
    parser.add_argument("--coelsa-pattern", default="*.txt", help="Patrón de archivos COELSA")
    parser.add_argument("--output-dir", default="./outputs", help="Directorio de salida")
    parser.add_argument("--output-excel", default="reporte_conciliacion_debin.xlsx", help="Nombre del archivo Excel de salida")
    parser.add_argument("--ftp-host", help="Host FTP")
    parser.add_argument("--ftp-user", help="Usuario FTP")
    parser.add_argument("--ftp-pass", help="Password FTP")
    parser.add_argument("--ftp-dir", help="Directorio FTP remoto")
    args = parser.parse_args()

    local_input_dir = Path(args.local_input_dir)
    output_dir = Path(args.output_dir)

    if args.ftp_host:
        required = [args.ftp_user, args.ftp_pass, args.ftp_dir]
        if not all(required):
            raise SystemExit("Si se usa FTP, también debe informar --ftp-user --ftp-pass --ftp-dir")
        download_from_ftp(args.ftp_host, args.ftp_user, args.ftp_pass, args.ftp_dir, local_input_dir, args.coelsa_pattern)

    coelsa_files = sorted(local_input_dir.glob(args.coelsa_pattern))
    if not coelsa_files:
        raise SystemExit(f"No se encontraron archivos COELSA con patrón {args.coelsa_pattern} en {local_input_dir}")

    core_csv = Path(args.core_csv)
    if not core_csv.exists():
        raise SystemExit(f"No existe el archivo Core CSV: {core_csv}")

    coelsa_movs = load_movements(coelsa_files, source="coelsa")
    core_movs = load_movements([core_csv], source="core")
    report = reconcile(coelsa_movs, core_movs)

    summary_rows = build_summary_rows(coelsa_movs, core_movs, report)
    excel_path = output_dir / args.output_excel
    write_excel_report(excel_path, summary_rows, report)

    counts = Counter(r["status"] for r in report)
    print(f"Conciliación generada: {excel_path}")
    print(
        " | ".join(
            [
                f"MATCH_ID_DEBIN={counts.get('MATCH_ID_DEBIN', 0)}",
                f"MATCH_COMNRO_CUENTA_IMPORTE={counts.get('MATCH_COMNRO_CUENTA_IMPORTE', 0)}",
                f"SOLO_COELSA={counts.get('SOLO_COELSA', 0)}",
                f"SOLO_CORE={counts.get('SOLO_CORE', 0)}",
            ]
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
